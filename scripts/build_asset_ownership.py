#!/usr/bin/env python3
"""
Genera api/_asset-ownership.js a partir del Excel de PROPIEDAD exportado
desde Metabase (hoja "Hoja 1" con columnas state, sku_id, nombre_sku,
serial, PROPIEDAD).

Ese Excel marca cada equipo del inventario como propiedad de BIA o de
Rentek. Los equipos Rentek NO se pueden vender, así que el proxy de
Metabase usa este listado para marcar cada equipo y el módulo de Venta
de Activos bloquea la cotización.

Uso:
    python3 scripts/build_asset_ownership.py ruta/al/PROPIEDAD.xlsx

El Excel no se versiona (contiene datos de inventario); solo se versiona
el .js generado. Para actualizar el listado, exporta de nuevo desde
Metabase y vuelve a correr este script.

El archivo vive dentro de api/ y con guion bajo al inicio a propósito:
Vercel no lo publica como endpoint y queda junto a la función que lo usa,
con la misma extensión .js que el resto de api/ (un .mjs fuera de api/
hacía fallar la función al desplegar).
"""

import re
import sys
from collections import defaultdict
from pathlib import Path

try:
    import openpyxl
except ImportError:
    sys.exit("Falta openpyxl. Instálalo con: pip install openpyxl")

SHEET = "Hoja 1"
OUT = Path(__file__).resolve().parent.parent / "api" / "_asset-ownership.js"


def norm(serial):
    """Clave de búsqueda: solo alfanuméricos, en mayúscula.

    Los seriales vienen con formatos muy distintos entre sistemas
    ('2025 / 103424', '2024-1004819096', '1B 21628'), así que se
    comparan ignorando separadores.
    """
    return re.sub(r"[^A-Za-z0-9]", "", str(serial or "")).upper()


def main():
    if len(sys.argv) < 2:
        sys.exit(f"Uso: python3 {sys.argv[0]} ruta/al/PROPIEDAD.xlsx")

    src = Path(sys.argv[1])
    if not src.exists():
        sys.exit(f"No existe el archivo: {src}")

    wb = openpyxl.load_workbook(src, read_only=True, data_only=True)
    if SHEET not in wb.sheetnames:
        sys.exit(f'El Excel no tiene la hoja "{SHEET}". Hojas: {wb.sheetnames}')

    rows = list(wb[SHEET].iter_rows(values_only=True))
    header = list(rows[0])
    i_serial = header.index("serial")
    i_prop = header.index("PROPIEDAD")

    owners = defaultdict(set)
    sin_serial = 0
    for row in rows[1:]:
        serial, prop = row[i_serial], row[i_prop]
        if serial in (None, "") or prop in (None, ""):
            sin_serial += 1
            continue
        key = norm(serial)
        if key:
            owners[key].add(str(prop).strip().upper())

    conflicts = {k: v for k, v in owners.items() if len(v) > 1}
    if conflicts:
        sys.exit(
            f"{len(conflicts)} seriales con propiedad contradictoria, "
            f"revisa el Excel: {list(conflicts.items())[:5]}"
        )

    rentek = sorted(k for k, v in owners.items() if "RENTEK" in v)
    bia = sorted(k for k, v in owners.items() if "BIA" in v)

    otros = sorted({p for v in owners.values() for p in v} - {"RENTEK", "BIA"})
    if otros:
        print(f"⚠️  Valores de PROPIEDAD no reconocidos (ignorados): {otros}")

    OUT.parent.mkdir(parents=True, exist_ok=True)
    OUT.write_text(
        '/* GENERADO POR scripts/build_asset_ownership.py — NO EDITAR A MANO.\n'
        f' * Fuente: {src.name} (hoja "{SHEET}")\n'
        f" * {len(rentek)} seriales Rentek · {len(bia)} seriales BIA\n"
        " *\n"
        " * Los seriales están normalizados (solo alfanuméricos, mayúscula)\n"
        " * porque el mismo equipo aparece con distintos separadores según\n"
        " * el sistema de origen. Usa normalizeSerial() para consultar.\n"
        " */\n\n"
        "// Un serial por línea: pesa bastante menos que un array JSON, se\n"
        "// convierte a Set al cargar y hace legibles los diffs de git.\n"
        f"const RENTEK_SERIALS = `{chr(10).join(rentek)}`;\n\n"
        f"const BIA_SERIALS = `{chr(10).join(bia)}`;\n\n"
        "const rentekSet = new Set(RENTEK_SERIALS.split('\\n'));\n"
        "const biaSet    = new Set(BIA_SERIALS.split('\\n'));\n\n"
        "export function normalizeSerial(serial) {\n"
        "  return String(serial || '').replace(/[^A-Za-z0-9]/g, '').toUpperCase();\n"
        "}\n\n"
        "/** 'RENTEK' | 'BIA' | null (null = sin serial o no está en el listado) */\n"
        "export function getAssetOwnership(serial) {\n"
        "  const key = normalizeSerial(serial);\n"
        "  if (!key) return null;\n"
        "  if (rentekSet.has(key)) return 'RENTEK';\n"
        "  if (biaSet.has(key)) return 'BIA';\n\n"
        "  // Algunos sistemas recortan los ceros a la izquierda del serial.\n"
        "  const sinCeros = key.replace(/^0+/, '');\n"
        "  if (sinCeros && sinCeros !== key) {\n"
        "    if (rentekSet.has(sinCeros)) return 'RENTEK';\n"
        "    if (biaSet.has(sinCeros)) return 'BIA';\n"
        "  }\n"
        "  return null;\n"
        "}\n\n"
        "export const ownershipStats = {\n"
        f"  rentek: {len(rentek)},\n"
        f"  bia: {len(bia)},\n"
        f"  total: {len(rentek) + len(bia)}\n"
        "};\n",
        encoding="utf-8",
    )

    kb = OUT.stat().st_size / 1024
    print(f"✓ {OUT.relative_to(OUT.parent.parent)} — {len(rentek)} Rentek, "
          f"{len(bia)} BIA ({kb:.0f} KB)")
    if sin_serial:
        print(f"  ({sin_serial} filas sin serial o sin propiedad, omitidas)")


if __name__ == "__main__":
    main()
